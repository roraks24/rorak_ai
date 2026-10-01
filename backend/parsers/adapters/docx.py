"""
DOCX document parser adapter.

Rorak AI V2.5 Step 3.1: Microsoft Word (.docx) document parser.
"""
import logging
import re
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional
from uuid import UUID

import docx

from backend.parsers.base import DocumentParser, ParsedDocument, StructuredBlock
from backend.parsers.registry import ParserError, UnsupportedFormatError


logger = logging.getLogger(__name__)


class DocxDocumentParser(DocumentParser):
    """
    Parser adapter for Microsoft Word (.docx) documents.
    Validates ZIP/OPC format, extracts paragraphs, headings, and tables,
    normalizes whitespace, and preserves hierarchical heading metadata.
    """

    SUPPORTED_EXTENSIONS = {"docx"}
    SUPPORTED_MIMES = {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/docx",
    }

    def supports(self, file_type: str) -> bool:
        """Return True if the file type matches DOCX extension or MIME type."""
        if not file_type:
            return False
        normalized = file_type.strip().lower().lstrip(".")
        return (
            normalized in self.SUPPORTED_EXTENSIONS
            or file_type.strip().lower() in self.SUPPORTED_MIMES
        )

    def _extract_heading_level(self, style_name: str) -> Optional[int]:
        """Extract numeric heading level from paragraph style name."""
        lower_style = style_name.strip().lower()
        if "title" in lower_style:
            return 1
        if "subtitle" in lower_style:
            return 2

        match = re.search(r"heading\s*(\d+)", lower_style)
        if match:
            try:
                return int(match.group(1))
            except ValueError:
                return 1

        return None

    def _normalize_text(self, text: str) -> str:
        """Normalize line breaks and collapse redundant horizontal whitespace."""
        if not text:
            return ""
        # Normalize carriage returns
        cleaned = text.replace("\r\n", "\n").replace("\r", "\n")
        # Collapse excessive inline horizontal spaces/tabs
        cleaned = re.sub(r"[ \t]+", " ", cleaned)
        return cleaned.strip()

    def parse(
        self,
        file_path: Path | str,
        document_id: Optional[UUID | str] = None,
        **kwargs: Any,
    ) -> ParsedDocument:
        """
        Parse a .docx document into a canonical ParsedDocument.
        """
        path_obj = Path(file_path)

        if not path_obj.exists():
            raise FileNotFoundError(f"DOCX file not found: {path_obj}")

        if not path_obj.is_file():
            raise ValueError(f"Path is not a file: {path_obj}")

        # Extension check
        if path_obj.suffix.lower() != ".docx":
            raise UnsupportedFormatError(
                f"File '{path_obj.name}' is not a .docx file (extension is '{path_obj.suffix}')."
            )

        # Validate that the file is a valid ZIP archive
        if not zipfile.is_zipfile(path_obj):
            raise ParserError(f"Invalid or corrupted DOCX archive (not a valid zip file): {path_obj.name}")

        # Verify Word document OPC entry exists
        try:
            with zipfile.ZipFile(path_obj) as zf:
                namelist = zf.namelist()
                if "word/document.xml" not in namelist:
                    raise ParserError(
                        f"Malformed DOCX: missing 'word/document.xml' in {path_obj.name}"
                    )
        except zipfile.BadZipFile as e:
            raise ParserError(f"Corrupted DOCX archive in {path_obj.name}: {e}")

        # Parse with python-docx
        try:
            doc = docx.Document(str(path_obj))
        except Exception as e:
            raise ParserError(f"Failed to read DOCX document {path_obj.name}: {e}")

        blocks: List[StructuredBlock] = []
        heading_stack: List[tuple[int, str]] = []  # (level, heading_text)
        current_heading: Optional[str] = None

        paragraph_idx = 0
        for p in doc.paragraphs:
            raw_text = p.text
            clean_text = self._normalize_text(raw_text)
            if not clean_text:
                continue

            style_name = p.style.name if p.style else "Normal"
            heading_level = self._extract_heading_level(style_name)

            if heading_level is not None:
                # Update hierarchical heading stack
                while heading_stack and heading_stack[-1][0] >= heading_level:
                    heading_stack.pop()
                heading_stack.append((heading_level, clean_text))
                current_heading = clean_text
                heading_path = " > ".join(h[1] for h in heading_stack)

                blocks.append(
                    StructuredBlock(
                        text=clean_text,
                        block_type="heading",
                        metadata={
                            "heading_level": heading_level,
                            "style": style_name,
                            "heading_text": clean_text,
                            "heading_path": heading_path,
                            "section": current_heading,
                        },
                    )
                )
            else:
                paragraph_idx += 1
                heading_path = " > ".join(h[1] for h in heading_stack) if heading_stack else ""
                block_meta: Dict[str, Any] = {
                    "style": style_name,
                    "paragraph_index": paragraph_idx,
                }
                if current_heading:
                    block_meta["section"] = current_heading
                if heading_path:
                    block_meta["heading_path"] = heading_path

                blocks.append(
                    StructuredBlock(
                        text=clean_text,
                        block_type="paragraph",
                        metadata=block_meta,
                    )
                )

        # Process tables if present
        table_idx = 0
        for table in doc.tables:
            table_rows: List[str] = []
            for row in table.rows:
                cells = [self._normalize_text(cell.text) for cell in row.cells]
                if any(cells):
                    table_rows.append(" | ".join(cells))

            if table_rows:
                table_idx += 1
                table_text = "\n".join(table_rows)
                heading_path = " > ".join(h[1] for h in heading_stack) if heading_stack else ""
                table_meta: Dict[str, Any] = {
                    "table_index": table_idx,
                    "row_count": len(table.rows),
                }
                if current_heading:
                    table_meta["section"] = current_heading
                if heading_path:
                    table_meta["heading_path"] = heading_path

                blocks.append(
                    StructuredBlock(
                        text=table_text,
                        block_type="table",
                        metadata=table_meta,
                    )
                )

        metadata: Dict[str, Any] = {
            "source_type": "docx",
            "file_size": path_obj.stat().st_size if path_obj.exists() else 0,
            "paragraph_count": paragraph_idx,
            "heading_count": len([b for b in blocks if b.block_type == "heading"]),
            "table_count": table_idx,
            "block_count": len(blocks),
        }

        return ParsedDocument(
            document_id=document_id,
            source_filename=path_obj.name,
            source_type="docx",
            structured_blocks=blocks,
            metadata=metadata,
        )
